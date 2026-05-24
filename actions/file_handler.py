"""
MARK-XX — File Analysis
Extract text from PDFs, images, code files, and plain text
for analysis by the LLM.
"""

import io
import os
from pathlib import Path
from typing import Optional, Tuple


SUPPORTED_EXTENSIONS = {
    "pdf": "pdf",
    "txt": "text",
    "md": "text",
    "py": "code",
    "js": "code",
    "ts": "code",
    "html": "code",
    "css": "code",
    "json": "code",
    "yaml": "code",
    "yml": "code",
    "xml": "code",
    "csv": "csv",
    "png": "image",
    "jpg": "image",
    "jpeg": "image",
    "gif": "image",
    "bmp": "image",
    "webp": "image",
    "docx": "docx",
    "xlsx": "xlsx",
}


def detect_file_type(path: str) -> str:
    ext = Path(path).suffix.lower().lstrip(".")
    return SUPPORTED_EXTENSIONS.get(ext, "unknown")


def extract_text(path: str) -> Tuple[str, str]:
    """
    Extract text from a file.
    Returns (extracted_text, file_type).
    """
    ftype = detect_file_type(path)
    p = Path(path)

    if not p.exists():
        return f"File not found: {path}", "error"

    if ftype == "pdf":
        return _extract_pdf(path), "pdf"
    elif ftype in ("text", "code", "csv"):
        return _extract_text_file(path), ftype
    elif ftype == "image":
        return "", "image"   # images are handled separately with vision
    elif ftype == "docx":
        return _extract_docx(path), "docx"
    elif ftype == "xlsx":
        return _extract_xlsx(path), "xlsx"
    else:
        # Try reading as plain text
        try:
            return p.read_text(encoding="utf-8", errors="replace"), "text"
        except Exception as e:
            return f"Cannot read file: {e}", "error"


def _extract_pdf(path: str) -> str:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(path)
        pages = []
        for i, page in enumerate(doc):
            text = page.get_text()
            if text.strip():
                pages.append(f"--- Page {i+1} ---\n{text.strip()}")
        doc.close()
        if pages:
            return "\n\n".join(pages)
        return "PDF appears to contain no extractable text (may be image-based)."
    except ImportError:
        return "PyMuPDF not installed. Run: pip install PyMuPDF"
    except Exception as e:
        return f"PDF extraction error: {e}"


def _extract_text_file(path: str) -> str:
    try:
        return Path(path).read_text(encoding="utf-8", errors="replace")
    except Exception as e:
        return f"Text extraction error: {e}"


def _extract_docx(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    except ImportError:
        return "python-docx not installed. Run: pip install python-docx"
    except Exception as e:
        return f"DOCX extraction error: {e}"


def _extract_xlsx(path: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        output = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            output.append(f"=== Sheet: {sheet_name} ===")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                output.append("\t".join(cells))
        return "\n".join(output)
    except ImportError:
        return "openpyxl not installed. Run: pip install openpyxl"
    except Exception as e:
        return f"XLSX extraction error: {e}"


def get_image_bytes(path: str) -> Optional[bytes]:
    """Read image file as raw bytes for vision analysis."""
    try:
        from PIL import Image
        img = Image.open(path)
        # Resize if very large
        if max(img.size) > 2048:
            img.thumbnail((2048, 2048))
        buf = io.BytesIO()
        fmt = img.format or "PNG"
        img.save(buf, format=fmt)
        return buf.getvalue()
    except Exception as e:
        print(f"Image read error: {e}")
        return None


def get_image_mime(path: str) -> str:
    ext = Path(path).suffix.lower()
    mime_map = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".gif": "image/gif",
        ".webp": "image/webp",
        ".bmp": "image/bmp",
    }
    return mime_map.get(ext, "image/png")


def summarize_file_info(path: str) -> str:
    """Return a short description of a file for context."""
    p = Path(path)
    ftype = detect_file_type(str(p))
    size = p.stat().st_size if p.exists() else 0
    size_str = f"{size:,} bytes" if size < 1024*1024 else f"{size/1024/1024:.1f} MB"
    return f"{p.name} ({ftype.upper()}, {size_str})"
